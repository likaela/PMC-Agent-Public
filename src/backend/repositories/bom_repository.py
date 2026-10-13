"""BOM Excel 数据读取。"""

from pathlib import Path

from openpyxl import load_workbook


DATA_FILE = Path(__file__).resolve().parents[3] / "data" / "bom.xlsx"
REQUIRED_COLUMNS = ("parent_material", "child_material", "qty_per")


def load_bom_data(file_path: Path = DATA_FILE) -> dict[str, dict[str, int]]:
    """读取 BOM 父子关系，返回 ``{parent_material: {child_material: qty_per}}``。"""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None or tuple(headers) != REQUIRED_COLUMNS:
            raise ValueError(f"BOM 文件必须包含列：{', '.join(REQUIRED_COLUMNS)}")

        result: dict[str, dict[str, int]] = {}
        for parent_material, child_material, qty_per in rows:
            if parent_material is None and child_material is None and qty_per is None:
                continue
            if (
                not isinstance(parent_material, str)
                or not isinstance(child_material, str)
                or not isinstance(qty_per, int)
                or isinstance(qty_per, bool)
                or qty_per <= 0
            ):
                raise ValueError("BOM 行必须包含产品、子物料和正整数单位用量")
            result.setdefault(parent_material, {})[child_material] = qty_per
        return result
    finally:
        workbook.close()
