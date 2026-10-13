"""受控库存工作簿适配器。

只读取已经确认的一期库存事实来源 ``库存表``，并保留行、库位和批次粒度。
其他工作表既不扫描明细，也不参与任何数量合并。
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import xlrd
from openpyxl import load_workbook


INVENTORY_SHEET_NAME = "库存表"
EXCLUDED_INVENTORY_SOURCES = (
    "电阻电容现库位",
    "2023年至2025年形成的存货-可以处理",
    "成品库存表",
    "料号查询表",
    "阻容库存表-2019",
)
INVENTORY_COLUMN_ALIASES = {
    "site": ("归属地",),
    "location": ("库位",),
    "material_code": ("物料代码",),
    "batch_no": ("批次号",),
    "on_hand_qty": ("库存",),
}


def _label(value: object) -> str:
    return re.sub(r"[\s　]+", "", str(value or "")).strip()


@dataclass(frozen=True)
class MappedInventoryRow:
    sheet_name: str
    source_row_number: int
    values: dict[str, object]


@dataclass(frozen=True)
class InventoryWorkbookResult:
    rows: tuple[MappedInventoryRow, ...]
    source_row_count: int
    detected_sheet: str | None
    excluded_sources: tuple[str, ...]


class InventoryWorkbookRepository:
    """以只读模式定位表头，仅流式读取匹配的库存明细列。"""

    header_scan_limit = 30
    scan_column_limit = 64

    def __init__(self, *, optional_column_aliases: dict[str, tuple[str, ...]] | None = None):
        self.optional_column_aliases = optional_column_aliases or {}
        if not isinstance(self.optional_column_aliases, dict):
            raise ValueError("可用量和预留量映射必须为字段字典")
        if set(self.optional_column_aliases) - {"available_qty", "reserved_qty"}:
            raise ValueError("仅支持显式声明的可用量和预留量映射")
        used_labels = {_label(alias) for aliases in INVENTORY_COLUMN_ALIASES.values() for alias in aliases}
        for aliases in self.optional_column_aliases.values():
            if (not isinstance(aliases, (list, tuple)) or not aliases
                    or any(not isinstance(alias, str) or not _label(alias) for alias in aliases)):
                raise ValueError("库存字段映射必须为明确的非空列名列表")
            labels = {_label(alias) for alias in aliases}
            if labels & used_labels:
                raise ValueError("库存字段映射不得把同一来源列解释为不同数量事实")
            used_labels.update(labels)

    def map_workbook(self, file_path: str | Path) -> InventoryWorkbookResult:
        source = Path(file_path)
        if source.suffix.lower() == ".xlsx":
            return self._map_xlsx(source)
        if source.suffix.lower() == ".xls":
            return self._map_xls(source)
        raise ValueError("库存导入仅支持 .xls 或 .xlsx 文件")

    def _map_xlsx(self, file_path: Path) -> InventoryWorkbookResult:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            if INVENTORY_SHEET_NAME not in workbook.sheetnames:
                return InventoryWorkbookResult((), 0, None, EXCLUDED_INVENTORY_SOURCES)
            sheet = workbook[INVENTORY_SHEET_NAME]
            rows = sheet.iter_rows(values_only=True, max_col=self.scan_column_limit)
            mapped, source_count = self._map_rows(INVENTORY_SHEET_NAME, rows)
            return InventoryWorkbookResult(
                tuple(mapped), source_count, INVENTORY_SHEET_NAME, EXCLUDED_INVENTORY_SOURCES
            )
        finally:
            workbook.close()

    def _map_xls(self, file_path: Path) -> InventoryWorkbookResult:
        workbook = xlrd.open_workbook(file_path, on_demand=True)
        try:
            if INVENTORY_SHEET_NAME not in workbook.sheet_names():
                return InventoryWorkbookResult((), 0, None, EXCLUDED_INVENTORY_SOURCES)
            sheet = workbook.sheet_by_name(INVENTORY_SHEET_NAME)
            rows = (
                tuple(sheet.row_values(index, end_colx=self.scan_column_limit))
                for index in range(sheet.nrows)
            )
            mapped, source_count = self._map_rows(INVENTORY_SHEET_NAME, rows)
            return InventoryWorkbookResult(
                tuple(mapped), source_count, INVENTORY_SHEET_NAME, EXCLUDED_INVENTORY_SOURCES
            )
        finally:
            if hasattr(workbook, "unload_sheet") and INVENTORY_SHEET_NAME in workbook.sheet_names():
                workbook.unload_sheet(INVENTORY_SHEET_NAME)
            workbook.release_resources()

    def _map_rows(
        self,
        sheet_name: str,
        rows: object,
    ) -> tuple[list[MappedInventoryRow], int]:
        alias_lookup = {
            _label(alias): field_name
            for field_name, aliases in {**INVENTORY_COLUMN_ALIASES, **self.optional_column_aliases}.items()
            for alias in aliases
        }
        header_indexes: dict[str, int] | None = None
        header_row_number = 0
        iterator = iter(rows)
        for row_number, row in enumerate(iterator, 1):
            indexes: dict[str, int] = {}
            for column_index, value in enumerate(row):
                field_name = alias_lookup.get(_label(value))
                if field_name and field_name not in indexes:
                    indexes[field_name] = column_index
            if set(INVENTORY_COLUMN_ALIASES).issubset(indexes):
                header_indexes = indexes
                header_row_number = row_number
                break
            if row_number >= self.header_scan_limit:
                break
        if header_indexes is None:
            return [], 0

        mapped: list[MappedInventoryRow] = []
        for source_row_number, row in enumerate(iterator, header_row_number + 1):
            values = {
                field_name: row[index] if index < len(row) else None
                for field_name, index in header_indexes.items()
            }
            if not any(value is not None and str(value).strip() for value in values.values()):
                continue
            mapped.append(MappedInventoryRow(sheet_name, source_row_number, values))
        return mapped, len(mapped)


__all__ = [
    "EXCLUDED_INVENTORY_SOURCES",
    "INVENTORY_SHEET_NAME",
    "InventoryWorkbookRepository",
    "InventoryWorkbookResult",
    "MappedInventoryRow",
]
