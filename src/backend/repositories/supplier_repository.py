"""真实供应商主数据的只读 Excel 适配器。"""

from pathlib import Path

from openpyxl import load_workbook

from repositories.normalization import normalize_header
from repositories.data_source import RAW_SUPPLIER_FILE, load_demo_data, resolve_data_source_mode


SUPPLIER_FILE = RAW_SUPPLIER_FILE
SUPPLIER_SHEET_NAME = "导出"


def load_supplier_master(file_path: Path | None = None) -> list[dict[str, str]]:
    """读取供应商主数据；不推断物料-供应商关系或采购参数。"""
    if file_path is None and resolve_data_source_mode() == "demo":
        return [dict(item) for item in load_demo_data()["supplier_master"]]

    workbook = load_workbook(file_path or SUPPLIER_FILE, read_only=True, data_only=True)
    try:
        sheet = workbook[SUPPLIER_SHEET_NAME]
        header_row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
        headers = {normalize_header(value): index for index, value in enumerate(header_row)}
        required_columns = {"公司编号", "公司名称"}
        if not required_columns.issubset(headers):
            raise ValueError("供应商表缺少公司编号或公司名称列")

        suppliers: list[dict[str, str]] = []
        for row in sheet.iter_rows(min_row=3, values_only=True):
            supplier_code = str(row[headers["公司编号"]] or "").strip()
            if not supplier_code:
                continue
            suppliers.append(
                {
                    "supplier_code": supplier_code,
                    "supplier_name": str(row[headers["公司名称"]] or "").strip(),
                    "industry_type": str(row[headers.get("行业类型", -1)] or "").strip() if "行业类型" in headers else "",
                    "supplier_grade": str(row[headers.get("供方等级", -1)] or "").strip() if "供方等级" in headers else "",
                    "remark": str(row[headers.get("备注", -1)] or "").strip() if "备注" in headers else "",
                }
            )
        return suppliers
    finally:
        workbook.close()
