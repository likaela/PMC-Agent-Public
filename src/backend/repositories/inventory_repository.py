"""库存 Excel 数据读取。"""

from pathlib import Path

from openpyxl import load_workbook


DATA_FILE = Path(__file__).resolve().parents[3] / "data" / "inventory.xlsx"
REQUIRED_COLUMNS = ("material", "stock", "reserved", "incoming", "safety_stock")


def load_inventory_data(file_path: Path = DATA_FILE) -> dict[str, dict[str, int]]:
    """读取库存主数据，返回按物料索引的库存记录。"""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None or tuple(headers) != REQUIRED_COLUMNS:
            raise ValueError(f"库存文件必须包含列：{', '.join(REQUIRED_COLUMNS)}")

        result: dict[str, dict[str, int]] = {}
        for material, stock, reserved, incoming, safety_stock in rows:
            if all(value is None for value in (material, stock, reserved, incoming, safety_stock)):
                continue
            numeric_values = (stock, reserved, incoming, safety_stock)
            if (
                not isinstance(material, str)
                or any(not isinstance(value, int) or isinstance(value, bool) or value < 0 for value in numeric_values)
            ):
                raise ValueError("库存行必须包含物料名称和非负整数数量")
            result[material] = {
                "stock": stock,
                "reserved": reserved,
                "incoming": incoming,
                "safety_stock": safety_stock,
            }
        return result
    finally:
        workbook.close()
