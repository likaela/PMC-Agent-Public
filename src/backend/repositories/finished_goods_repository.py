"""真实成品库存表的只读 Excel 适配器；一期不参与 MRP。"""

from pathlib import Path

from openpyxl import load_workbook

from repositories.normalization import normalize_header, normalize_material_code, normalize_nonnegative_int


RAW_DATA_DIR = Path(__file__).resolve().parents[3] / "data" / "raw"
FINISHED_GOODS_FILE = RAW_DATA_DIR / "2026年库存表+料号查询表及成品库存表(1).xlsx"
FINISHED_GOODS_SHEET_NAME = "成品库存表"


def load_finished_goods_details(file_path: Path = FINISHED_GOODS_FILE) -> list[dict[str, object]]:
    """读取成品库存明细，保留供后续阶段使用。"""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook[FINISHED_GOODS_SHEET_NAME]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {normalize_header(value): index for index, value in enumerate(header_row)}
        required_columns = {"归属地", "位置", "物料代码", "批次", "库存"}
        if not required_columns.issubset(headers):
            raise ValueError("成品库存表缺少必需列")

        result: list[dict[str, object]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            material_code = normalize_material_code(row[headers["物料代码"]])
            if material_code:
                result.append(
                    {
                        "material_code": material_code,
                        "site": str(row[headers["归属地"]] or "").strip(),
                        "location": str(row[headers["位置"]] or "").strip(),
                        "batch": str(row[headers["批次"]] or "").strip(),
                        "stock": normalize_nonnegative_int(row[headers["库存"]], "成品库存"),
                        "source_sheet": FINISHED_GOODS_SHEET_NAME,
                    }
                )
        return result
    finally:
        workbook.close()
