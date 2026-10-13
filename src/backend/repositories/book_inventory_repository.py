"""Read-only, explicitly labelled Excel observations, never published availability.

Each physical row contributes once to a *recorded* total. Identical observation
rows are retained and flagged, not silently deduplicated or used for decisions.
No old Demo exclusions, transaction columns, reserved defaults or source merging.
"""
from __future__ import annotations

import copy
import hashlib
import os
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from threading import RLock

from openpyxl import load_workbook
from repositories.normalization import normalize_material_code


class BookInventoryRepository:
    _cache: dict = {}
    _lock = RLock()

    def __init__(self, path: Path, *, sheet_name: str = "库存表", unit_confirmations=None, quantity_confirmations=None):
        self.path = Path(path)
        self.sheet_name = sheet_name
        from repositories.inventory_unit_confirmation_repository import InventoryUnitConfirmationRepository
        self.unit_confirmations = unit_confirmations or InventoryUnitConfirmationRepository()
        from repositories.inventory_unit_confirmation_repository import InventoryQuantityConfirmationRepository
        self.quantity_confirmations = quantity_confirmations or InventoryQuantityConfirmationRepository()

    def _confirmed_units(self, payload):
        # Apply *after* the content cache: a later human confirmation must be read
        # even when the original Excel has not changed. Never mutate cached rows.
        payload = copy.deepcopy(payload)
        confirmation = self.unit_confirmations.for_source(self.path, payload["content_sha256"], self.sheet_name)
        payload["unit_confirmation"] = confirmation
        payload["quantity_confirmation"] = self.quantity_confirmations.for_source(
            self.path, payload["content_sha256"], self.sheet_name)
        for details in payload["records"].values():
            for row in details:
                raw_unit = row["inventory_unit"]
                row["raw_inventory_unit"] = raw_unit
                if confirmation and raw_unit not in (None, confirmation["unit"]):
                    row["inventory_unit"] = None
                    payload["row_errors"].append({"row_number": row["provenance"]["row_number"],
                        "material_code": row["material_code"], "reason": "inventory_unit_conflict"})
                elif confirmation:
                    row["inventory_unit"] = confirmation["unit"]
                    row["provenance"]["unit_confirmation"] = copy.deepcopy(confirmation)
        return payload

    def load(self) -> dict:
        if os.getenv("PMC_TRIAL_MODE") == "1":
            raise PermissionError("Private Excel observations are disabled in trial")
        digest = hashlib.sha256(self.path.read_bytes()).hexdigest()
        key = (str(self.path.resolve()), self.sheet_name, digest)
        with self._lock:
            if key in self._cache:
                return self._confirmed_units(self._cache[key])
            workbook = load_workbook(self.path, read_only=True, data_only=True)
            try:
                sheet = workbook[self.sheet_name]
                # Never scan the hundreds of movement/date columns as stock facts.
                header = next(sheet.iter_rows(max_row=1, values_only=True))
                headers = {str(v).strip(): i for i, v in enumerate(header) if v is not None}
                required = ("物料代码", "库存", "库位", "批次号", "归属地")
                if not set(required) <= headers.keys():
                    raise ValueError("Main inventory mapping does not match source")
                selected = (*required, "名称", "单位")
                last = max(headers[k] for k in selected if k in headers) + 1
                records, errors, signatures = defaultdict(list), [], Counter()
                for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=last, values_only=True), 2):
                    def value(field):
                        return row[headers[field]] if field in headers else None
                    if all(v is None for v in row):
                        continue
                    code = normalize_material_code(value("物料代码"))
                    if not code:
                        errors.append({"row_number": row_number, "reason": "material_code_missing"})
                        continue
                    raw = value("库存")
                    try:
                        if raw is None or isinstance(raw, bool):
                            raise ValueError()
                        quantity = Decimal(str(raw))
                        if not quantity.is_finite() or quantity < 0:
                            raise ValueError()
                    except (InvalidOperation, ValueError):
                        quantity = None
                        errors.append({"row_number": row_number, "material_code": code, "reason": "stock_invalid"})
                    def label(field):
                        v = value(field)
                        return str(v).strip() if v is not None and str(v).strip() else None
                    detail = dict(material_code=code, name=label("名称"), recorded_stock_qty=None if quantity is None else format(quantity, "f"),
                        location=label("库位"), batch_no=label("批次号"), site=label("归属地"), inventory_unit=label("单位"),
                        provenance=dict(source_file=self.path.name, content_sha256=digest, sheet_name=self.sheet_name,
                            row_number=row_number, source_kind="excel_observation", snapshot_at=None, import_batch_id=None))
                    signature = tuple(detail[k] for k in ("material_code", "recorded_stock_qty", "location", "batch_no", "site"))
                    signatures[signature] += 1
                    detail["signature"] = signature
                    records[code].append(detail)
                for details in records.values():
                    for detail in details:
                        detail["duplicate_observation"] = signatures[detail.pop("signature")] > 1
                payload = dict(records=dict(records), source_file=self.path.name, content_sha256=digest,
                    sheet_name=self.sheet_name, record_count=sum(map(len, records.values())), unique_material_count=len(records),
                    row_errors=errors, duplicate_observation_count=sum(n - 1 for n in signatures.values() if n > 1),
                    excluded_sheets=[s for s in workbook.sheetnames if s != self.sheet_name])
            finally:
                workbook.close()
            if hashlib.sha256(self.path.read_bytes()).hexdigest() != digest:
                raise ValueError("Source changed during read; retry against a stable workbook")
            # Bounded content-addressed cache; source changes never reuse old observations.
            if len(self._cache) >= 4:
                self._cache.clear()
            self._cache[key] = payload
            return self._confirmed_units(payload)
