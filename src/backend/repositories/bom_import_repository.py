"""通用 BOM Excel 流式源文件适配器。

适配器只扫描每个工作表开头有限行来识别候选表头；命中后才逐行读取 BOM 明细。
它不依赖固定文件名、工作表名或产品，也不会把整个工作簿构造成内存二维表。
"""

from __future__ import annotations

import json
import re
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

import xlrd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_3588M_MAPPING = PROJECT_ROOT / "import_mappings" / "3588m_costing_v1.json"


def _normalized_label(value: object) -> str:
    return re.sub(r"[\s　]+", "", str(value or "")).strip()


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


@dataclass(frozen=True)
class BomImportMapping:
    mapping_name: str
    metadata_aliases: dict[str, tuple[str, ...]]
    column_aliases: dict[str, tuple[str, ...]]
    required_columns: tuple[str, ...]
    pending_when_absent: tuple[str, ...]
    stop_row_markers: tuple[str, ...]
    header_scan_limit: int
    total_price_tolerance: str
    price_policy: str = "required"
    publish_parent_item: bool = False
    included_sheets: tuple[str, ...] = ()
    unit_policy: str = "required"
    publication_scope: str = "full"


@dataclass(frozen=True)
class MappedBomRow:
    sheet_name: str
    source_row_number: int
    parent_material_code: str | None
    version: str | None
    values: dict[str, object]
    metadata_row_number: int | None = None


@dataclass(frozen=True)
class WorkbookMappingResult:
    rows: tuple[MappedBomRow, ...]
    detected_sheets: tuple[str, ...]
    sheets_without_metadata: tuple[str, ...]
    scanned_row_count: int


def load_bom_import_mapping(file_path: str | Path = DEFAULT_3588M_MAPPING) -> BomImportMapping:
    mapping_path = Path(file_path)
    payload = json.loads(mapping_path.read_text(encoding="utf-8"))
    header_scan_limit = int(payload.get("header_scan_limit", 100))
    if header_scan_limit <= 0:
        raise ValueError("header_scan_limit 必须大于 0")
    price_policy = payload.get("price_policy", "required")
    if price_policy not in {"required", "optional"}:
        raise ValueError("price_policy 必须为 required 或 optional")
    unit_policy = payload.get("unit_policy", "required")
    if unit_policy not in {"required", "optional_for_structure"}:
        raise ValueError("unit_policy 必须为 required 或 optional_for_structure")
    publish_parent = payload.get("publish_parent_item", False)
    if not isinstance(publish_parent, bool):
        raise ValueError("publish_parent_item 必须为布尔值")
    scope = payload.get("publication_scope", "full")
    if scope not in {"full", "structure"}:
        raise ValueError("publication_scope 必须为 full 或 structure")
    if scope == "structure" and (not publish_parent or price_policy != "optional" or unit_policy != "optional_for_structure"):
        raise ValueError("结构发布必须显式允许未知单位/价格并发布父项")
    return BomImportMapping(
        mapping_name=str(payload["mapping_name"]),
        metadata_aliases={
            name: tuple(str(alias) for alias in definition["aliases"])
            for name, definition in payload["metadata"].items()
        },
        column_aliases={
            name: tuple(str(alias) for alias in definition["aliases"])
            for name, definition in payload["columns"].items()
        },
        required_columns=tuple(str(value) for value in payload["required_columns"]),
        pending_when_absent=tuple(str(value) for value in payload["pending_when_absent"]),
        stop_row_markers=tuple(str(value) for value in payload.get("stop_row_markers", ())),
        header_scan_limit=header_scan_limit,
        total_price_tolerance=str(payload["validation"]["total_price_tolerance"]),
        price_policy=price_policy,
        publish_parent_item=publish_parent,
        included_sheets=tuple(payload.get("included_sheets", ())),
        unit_policy=unit_policy,
        publication_scope=scope,
    )


class BomWorkbookRepository:
    """按工作表流式识别并读取旧版/新版 Excel BOM。"""

    def map_workbook(
        self,
        file_path: str | Path,
        mapping: BomImportMapping,
    ) -> WorkbookMappingResult:
        source_path = Path(file_path)
        suffix = source_path.suffix.lower()
        if suffix == ".xls":
            return self._map_xls(source_path, mapping)
        if suffix == ".xlsx":
            return self._map_xlsx(source_path, mapping)
        raise ValueError("BOM 导入仅支持 .xls 或 .xlsx 文件")

    def _map_xls(self, file_path: Path, mapping: BomImportMapping) -> WorkbookMappingResult:
        workbook = xlrd.open_workbook(file_path, on_demand=True)
        mapped_rows: list[MappedBomRow] = []
        detected_sheets: list[str] = []
        sheets_without_metadata: list[str] = []
        scanned_row_count = 0
        try:
            for sheet_name in workbook.sheet_names():
                if mapping.included_sheets and sheet_name not in mapping.included_sheets:
                    continue
                sheet = workbook.sheet_by_name(sheet_name)
                rows = (tuple(sheet.row_values(index)) for index in range(sheet.nrows))
                sheet_rows, detected, missing_metadata, scanned = self._map_sheet_rows(
                    sheet_name,
                    rows,
                    mapping,
                )
                mapped_rows.extend(sheet_rows)
                scanned_row_count += scanned
                if detected:
                    detected_sheets.append(sheet_name)
                if missing_metadata:
                    sheets_without_metadata.append(sheet_name)
                if hasattr(workbook, "unload_sheet"):
                    workbook.unload_sheet(sheet_name)
        finally:
            workbook.release_resources()
        return WorkbookMappingResult(
            rows=tuple(mapped_rows),
            detected_sheets=tuple(detected_sheets),
            sheets_without_metadata=tuple(sheets_without_metadata),
            scanned_row_count=scanned_row_count,
        )

    def _map_xlsx(self, file_path: Path, mapping: BomImportMapping) -> WorkbookMappingResult:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        mapped_rows: list[MappedBomRow] = []
        detected_sheets: list[str] = []
        sheets_without_metadata: list[str] = []
        scanned_row_count = 0
        try:
            for sheet in workbook.worksheets:
                if mapping.included_sheets and sheet.title not in mapping.included_sheets:
                    continue
                rows = (tuple(row) for row in sheet.iter_rows(values_only=True))
                sheet_rows, detected, missing_metadata, scanned = self._map_sheet_rows(
                    sheet.title,
                    rows,
                    mapping,
                )
                mapped_rows.extend(sheet_rows)
                scanned_row_count += scanned
                if detected:
                    detected_sheets.append(sheet.title)
                if missing_metadata:
                    sheets_without_metadata.append(sheet.title)
        finally:
            workbook.close()
        return WorkbookMappingResult(
            rows=tuple(mapped_rows),
            detected_sheets=tuple(detected_sheets),
            sheets_without_metadata=tuple(sheets_without_metadata),
            scanned_row_count=scanned_row_count,
        )

    def _map_sheet_rows(
        self,
        sheet_name: str,
        rows: Iterator[tuple[object, ...]],
        mapping: BomImportMapping,
    ) -> tuple[list[MappedBomRow], bool, bool, int]:
        metadata = {field_name: None for field_name in mapping.metadata_aliases}
        header_indexes: dict[str, int] | None = None
        header_row_number: int | None = None
        metadata_row_number: int | None = None
        scanned = 0

        for zero_based_index, row in enumerate(rows):
            scanned += 1
            before = dict(metadata)
            self._update_metadata(metadata, row, mapping)
            if metadata_row_number is None and metadata != before:
                metadata_row_number = zero_based_index + 1
            candidate_indexes = self._column_indexes(row, mapping)
            if all(field_name in candidate_indexes for field_name in mapping.required_columns):
                header_indexes = candidate_indexes
                header_row_number = zero_based_index + 1
                break
            if scanned >= mapping.header_scan_limit:
                return [], False, False, scanned

        if header_indexes is None or header_row_number is None:
            return [], False, False, scanned

        missing_metadata = any(not metadata.get(field_name) for field_name in mapping.metadata_aliases)
        stop_markers = {_normalized_label(marker) for marker in mapping.stop_row_markers}
        mapped_rows: list[MappedBomRow] = []
        for source_row_number, row in enumerate(rows, header_row_number + 1):
            scanned += 1
            values = {
                field_name: (row[column_index] if column_index < len(row) else None)
                for field_name, column_index in header_indexes.items()
            }
            if not any(value is not None and str(value).strip() for value in values.values()):
                continue
            if any(
                _normalized_label(value) in stop_markers
                for value in values.values()
                if _text(value) is not None
            ):
                break
            mapped_rows.append(
                MappedBomRow(
                    sheet_name=sheet_name,
                    source_row_number=source_row_number,
                    parent_material_code=metadata.get("parent_material_code"),
                    version=metadata.get("version"),
                    values=values,
                    metadata_row_number=metadata_row_number,
                )
            )
        return mapped_rows, True, missing_metadata, scanned

    @staticmethod
    def _column_indexes(
        row: tuple[object, ...],
        mapping: BomImportMapping,
    ) -> dict[str, int]:
        alias_lookup = {
            _normalized_label(alias): field_name
            for field_name, aliases in mapping.column_aliases.items()
            for alias in aliases
        }
        indexes: dict[str, int] = {}
        for column_index, cell_value in enumerate(row):
            field_name = alias_lookup.get(_normalized_label(cell_value))
            if field_name is not None and field_name not in indexes:
                indexes[field_name] = column_index
        return indexes

    @staticmethod
    def _update_metadata(
        result: dict[str, str | None],
        row: tuple[object, ...],
        mapping: BomImportMapping,
    ) -> None:
        for cell_index, cell_value in enumerate(row):
            cell_text = _text(cell_value)
            if not cell_text:
                continue
            for field_name, aliases in mapping.metadata_aliases.items():
                for alias in aliases:
                    if _normalized_label(cell_text) == _normalized_label(alias):
                        adjacent = next(
                            (
                                _text(value)
                                for value in row[cell_index + 1 :]
                                if _text(value) is not None
                            ),
                            None,
                        )
                        if adjacent:
                            if result[field_name] is not None and result[field_name] != adjacent:
                                raise ValueError("BOM 元数据存在冲突")
                            result[field_name] = adjacent
                            break
                    matches = re.findall(
                        rf"{re.escape(alias)}\s*[:：]\s*([^\s,，;；]+)",
                        cell_text,
                        flags=re.IGNORECASE,
                    )
                    if len(set(matches)) > 1:
                        raise ValueError("BOM 元数据存在冲突")
                    if matches:
                        value = matches[0].strip()
                        if result[field_name] is not None and result[field_name] != value:
                            raise ValueError("BOM 元数据存在冲突")
                        result[field_name] = value
                        break


__all__ = [
    "BomImportMapping",
    "BomWorkbookRepository",
    "DEFAULT_3588M_MAPPING",
    "MappedBomRow",
    "WorkbookMappingResult",
    "load_bom_import_mapping",
]
