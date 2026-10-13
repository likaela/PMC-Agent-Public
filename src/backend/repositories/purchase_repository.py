"""采购参数 Excel 数据读取。"""

from pathlib import Path

from openpyxl import load_workbook


DATA_FILE = Path(__file__).resolve().parents[3] / "data" / "purchase.xlsx"
REQUIRED_COLUMNS = ("material", "safety_stock", "lead_time_days")


def load_purchase_data(file_path: Path = DATA_FILE) -> dict[str, dict[str, int]]:
    """读取采购参数，返回按物料索引的安全库存和采购周期。"""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None or tuple(headers) != REQUIRED_COLUMNS:
            raise ValueError(f"采购文件必须包含列：{', '.join(REQUIRED_COLUMNS)}")

        result: dict[str, dict[str, int]] = {}
        for material, safety_stock, lead_time_days in rows:
            if all(value is None for value in (material, safety_stock, lead_time_days)):
                continue
            if (
                not isinstance(material, str)
                or not isinstance(safety_stock, int)
                or not isinstance(lead_time_days, int)
                or isinstance(safety_stock, bool)
                or isinstance(lead_time_days, bool)
                or safety_stock < 0
                or lead_time_days < 0
            ):
                raise ValueError("采购行必须包含物料名称和非负整数参数")
            result[material] = {
                "safety_stock": safety_stock,
                "lead_time_days": lead_time_days,
            }
        return result
    finally:
        workbook.close()
