"""库存业务确认表的只读 Excel 适配器。"""

from pathlib import Path

from openpyxl import load_workbook

from repositories.normalization import normalize_header, normalize_material_code, normalize_nonnegative_number


DISPOSITION_FILE = (
    Path(__file__).resolve().parents[3] / "data" / "templates" / "material_inventory_disposition.xlsx"
)
DISPOSITION_SHEET_NAME = "Dispositions"
ALLOWED_DISPOSITIONS = {
    "stock_zero",
    "non_stock_item",
    "purchase_to_order",
    "customer_supplied",
    "outsourced",
    "bom_obsolete",
    "inventory_missing",
    "pending_confirmation",
}
REQUIRED_COLUMNS = {
    "material_code",
    "disposition",
    "confirmed_stock",
    "include_in_mrp",
    "reason",
    "confirmed_by",
    "remark",
}


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def load_inventory_dispositions(file_path: Path = DISPOSITION_FILE) -> dict[str, dict[str, object]]:
    """读取人工确认记录；空 disposition 行视为尚未确认，不做任何状态推断。"""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook[DISPOSITION_SHEET_NAME]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {normalize_header(value): index for index, value in enumerate(header_row)}
        if not REQUIRED_COLUMNS.issubset(headers):
            raise ValueError("库存业务确认表缺少必需列")

        result: dict[str, dict[str, object]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            material_code = normalize_material_code(row[headers["material_code"]])
            disposition = str(row[headers["disposition"]] or "").strip()
            if not material_code or not disposition:
                continue
            if disposition not in ALLOWED_DISPOSITIONS:
                raise ValueError(f"料号 {material_code} 的 disposition 不受支持：{disposition}")
            if material_code in result:
                raise ValueError(f"库存业务确认表存在重复料号：{material_code}")

            include_value = row[headers["include_in_mrp"]]
            if not isinstance(include_value, bool):
                raise ValueError(f"料号 {material_code} 必须明确填写 include_in_mrp 为 TRUE 或 FALSE")
            stock_value = row[headers["confirmed_stock"]]
            confirmed_stock = None if _is_blank(stock_value) else normalize_nonnegative_number(stock_value, "confirmed_stock")
            if disposition == "stock_zero":
                if confirmed_stock != 0 or not include_value:
                    raise ValueError(
                        f"料号 {material_code} 的 stock_zero 必须同时填写 confirmed_stock=0 和 include_in_mrp=TRUE"
                    )

            result[material_code] = {
                "material_code": material_code,
                "disposition": disposition,
                "confirmed_stock": confirmed_stock,
                "include_in_mrp": include_value,
                "reason": str(row[headers["reason"]] or "").strip(),
                "confirmed_by": str(row[headers["confirmed_by"]] or "").strip(),
                "remark": str(row[headers["remark"]] or "").strip(),
            }
        return result
    finally:
        workbook.close()
