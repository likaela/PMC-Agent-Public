"""真实料号查询表的只读 Excel 适配器。"""

from pathlib import Path

from openpyxl import load_workbook

from repositories.normalization import normalize_header, normalize_material_code
from repositories.data_source import RAW_INVENTORY_FILE, load_demo_data, resolve_data_source_mode


MASTER_FILE = RAW_INVENTORY_FILE
MASTER_SHEET_NAME = "料号查询表"


def load_material_master_records(file_path: Path | None = None) -> list[dict[str, str]]:
    """读取料号主数据，不参与 BOM 或库存数量计算。"""
    if file_path is None and resolve_data_source_mode() == "demo":
        return [
            {
                "material_code": str(item["child_material"]),
                "name": f"演示物料 {index:03d}",
                "specification": "SYNTHETIC",
                "package": "DEMO",
                "quality_grade": "演示",
                "brand_or_manufacturer": "无真实厂商",
                "remark": "纯合成公开演示数据",
            }
            for index, item in enumerate(load_demo_data()["bom_relationships"], start=1)
        ]

    workbook = load_workbook(file_path or MASTER_FILE, read_only=True, data_only=True)
    try:
        sheet = workbook[MASTER_SHEET_NAME]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {normalize_header(value): index for index, value in enumerate(header_row)}
        required_columns = {"料号", "名称", "规格型号", "封装", "质量等级", "品牌/厂商", "备注"}
        if not required_columns.issubset(headers):
            raise ValueError("料号查询表缺少一期必需列")

        records: list[dict[str, str]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            material_code = normalize_material_code(row[headers["料号"]])
            if not material_code:
                continue
            records.append(
                {
                    "material_code": material_code,
                    "name": str(row[headers["名称"]] or "").strip(),
                    "specification": str(row[headers["规格型号"]] or "").strip(),
                    "package": str(row[headers["封装"]] or "").strip(),
                    "quality_grade": str(row[headers["质量等级"]] or "").strip(),
                    "brand_or_manufacturer": str(row[headers["品牌/厂商"]] or "").strip(),
                    "remark": str(row[headers["备注"]] or "").strip(),
                }
            )
        return records
    finally:
        workbook.close()


def index_material_master(records: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    """按料号索引主数据；重复料号由质量报告单独列示。"""
    return {record["material_code"]: record for record in records}
